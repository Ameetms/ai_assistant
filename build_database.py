#!/usr/bin/env python3
"""
build_database.py   (Step 2 of the project: the numbers)
=========================================================
Reads the SEC "Financial_Report" Excel files you downloaded (one per annual report)
and builds a clean SQL database of Apple's annual numbers.

Each annual report shows three fiscal years, so the files overlap. The script
 * keeps the value from the NEWEST report for each year,
 * compares the overlapping values and tells you if two reports disagree,
 * checks the accounting (categories add up to net sales, regions add up, ...).

Usage (from the project folder)
  python build_database.py --excel excel
  python build_database.py --excel "D:\\ai_assistant_project\\excel" --out database

Output folder (default: database/)
  apple.db                   SQLite database (open with DB Browser for SQLite, or Python)
  Apple_Annual_Tables.xlsx   the same tables in Excel
  tables/*.csv               the same tables as CSV
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
DATE_RE = re.compile(r"([A-Za-z]{3,5})\.? (\d{1,2}), (\d{4})")

SHEET_PATTERNS = {
    "income": re.compile(r"^CONSOLIDATED STATEMENTS OF OPERATIONS", re.I),
    "balance": re.compile(r"^CONSOLIDATED BALANCE SHEETS\s+-", re.I),
    "cashflow": re.compile(r"^CONSOLIDATED STATEMENTS OF CASH FLOWS", re.I),
    "category": re.compile(r"Net Sales Disaggregated by Significant Products and Services", re.I),
    "segment": re.compile(r"Information by Reportable Segment", re.I),
    "country": re.compile(r"Segment Information and Geographic Data - Net Sales \(Details\)", re.I),
}

INCOME_LABELS = {
    "net sales": "net_sales",
    "cost of sales": "cost_of_sales",
    "gross margin": "gross_margin",
    "research and development": "research_and_development",
    "selling, general and administrative": "selling_general_admin",
    "total operating expenses": "total_operating_expenses",
    "operating income": "operating_income",
    "other income/(expense), net": "other_income_expense_net",
    "income before provision for income taxes": "income_before_taxes",
    "provision for income taxes": "income_tax_provision",
    "net income": "net_income",
    "basic (in dollars per share)": "eps_basic",
    "diluted (in dollars per share)": "eps_diluted",
    "basic (in shares)": "shares_basic_thousands",
    "diluted (in shares)": "shares_diluted_thousands",
}
INCOME_ORDER = [
    "net_sales", "products_net_sales", "services_net_sales", "cost_of_sales", "products_cost_of_sales",
    "services_cost_of_sales", "gross_margin", "research_and_development", "selling_general_admin",
    "total_operating_expenses", "operating_income", "other_income_expense_net", "income_before_taxes",
    "income_tax_provision", "net_income", "eps_basic", "eps_diluted", "shares_basic_thousands",
    "shares_diluted_thousands",
]
CATEGORIES = ["iPhone", "Mac", "iPad", "Wearables, Home and Accessories", "Services"]
REGIONS = ["Americas", "Europe", "Greater China", "Japan", "Rest of Asia Pacific"]
COUNTRIES = {"u.s.": "United States", "us": "United States", "united states": "United States",
             "china": "China", "other countries": "Other countries"}


# ---------------------------------------------------------------------------
# reading the workbooks
# ---------------------------------------------------------------------------
def fiscal_year_from_date(text):
    m = DATE_RE.search(str(text))
    if not m:
        return None
    mon = MONTHS.get(m.group(1)[:3].lower())
    if not mon:
        return None
    year = int(m.group(3))
    return year + 1 if mon >= 10 else year  # Apple's year ends in September


def find_header(rows):
    best_i, best_cols = None, {}
    for i, r in enumerate(rows[:4]):
        cols = {}
        for j, c in enumerate(r):
            if j == 0 or c is None:
                continue
            fy = fiscal_year_from_date(c)
            if fy:
                cols[j] = fy
        if len(cols) > len(best_cols):
            best_i, best_cols = i, cols
    return best_i, best_cols


def parse_rows(rows):
    """Turn a sheet into [{section, label, values:{fiscal_year: number}}]."""
    hdr, cols = find_header(rows)
    if hdr is None:
        return []
    out, section = [], None
    for r in rows[hdr + 1:]:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        label = str(r[0]).strip()
        nums = {}
        for c, fy in cols.items():
            v = r[c] if c < len(r) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                nums[fy] = float(v)
        if not nums:
            if re.search(r"\[(Line Items|Abstract|Axis|Member|Table|Domain)\]", label):
                continue
            section = label.rstrip(":").strip()
            continue
        out.append({"section": section, "label": label, "values": nums})
    return out


def open_workbook_sheets(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found = {}
    for name in wb.sheetnames:
        ws = wb[name]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), (None,))
        title = str(first[0] or "")
        for key, pat in SHEET_PATTERNS.items():
            if key not in found and pat.search(title):
                found[key] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return found


# ---------------------------------------------------------------------------
class Collector:
    """Gathers (key -> value) from several files and keeps the newest report's value."""

    def __init__(self):
        self.data = defaultdict(dict)  # table -> key -> list[(source_year, value, file)]

    def add(self, table, key, value, source_year, filename):
        self.data[table].setdefault(key, []).append((source_year, value, filename))

    def final(self, table):
        return {k: max(v, key=lambda x: x[0])[1] for k, v in self.data[table].items()}

    def disagreements(self, tol_default=0.5):
        rows = []
        for table, d in self.data.items():
            for key, vals in d.items():
                if len(vals) < 2:
                    continue
                tol = 0.006 if "eps" in str(key).lower() else tol_default
                lo, hi = min(v[1] for v in vals), max(v[1] for v in vals)
                if hi - lo > tol:
                    rows.append({"table": table, "item": " / ".join(map(str, key)),
                                 "values_by_report": "; ".join("FY%d report: %s" % (v[0], v[1]) for v in sorted(vals))})
        return rows


def process_file(path: Path, col: Collector, log: list):
    sheets = open_workbook_sheets(path)
    # which report is this? the newest fiscal year in the income statement header
    inc = parse_rows(sheets.get("income", []))
    if not inc:
        log.append("!! %s: could not find the income statement sheet - skipped" % path.name)
        return None
    src = max(fy for e in inc for fy in e["values"])
    for key in SHEET_PATTERNS:
        if key not in sheets:
            log.append("note: %s: no sheet found for '%s' (that table will be missing for this report)" % (path.name, key))

    # --- income statement (wide)
    unmapped = set()
    for e in inc:
        base = INCOME_LABELS.get(e["label"].lower())
        if base is None:
            unmapped.add(e["label"])
            continue
        sec = (e["section"] or "").lower()
        key = "%s_%s" % (sec, base) if base in ("net_sales", "cost_of_sales") and sec in ("products", "services") else base
        for fy, v in e["values"].items():
            col.add("income_statement", (fy, key), v, src, path.name)
    if unmapped:
        log.append("note: %s: income lines not mapped to a column (kept in statement_lines): %s" % (path.name, sorted(unmapped)))

    # --- three statements in long form (nothing is lost)
    for stmt, key in (("Income statement", "income"), ("Balance sheet", "balance"), ("Cash flow statement", "cashflow")):
        for e in parse_rows(sheets.get(key, [])):
            for fy, v in e["values"].items():
                col.add("statement_lines", (stmt, fy, e["section"] or "", e["label"]), v, src, path.name)

    # --- net sales by category
    for e in parse_rows(sheets.get("category", [])):
        if e["label"].lower() == "net sales" and e["section"] in CATEGORIES:
            for fy, v in e["values"].items():
                col.add("net_sales_by_category", (fy, e["section"]), v, src, path.name)

    # --- regional segments
    for e in parse_rows(sheets.get("segment", [])):
        if e["section"] in REGIONS and e["label"].lower() in ("net sales", "operating income"):
            metric = e["label"].lower().replace(" ", "_")
            for fy, v in e["values"].items():
                col.add("segment_results", (fy, e["section"], metric), v, src, path.name)

    # --- countries
    for e in parse_rows(sheets.get("country", [])):
        name = COUNTRIES.get((e["section"] or "").lower())
        if name and e["label"].lower() == "net sales":
            for fy, v in e["values"].items():
                col.add("net_sales_by_country", (fy, name), v, src, path.name)
    return src, sorted({fy for e in inc for fy in e["values"]})


# ---------------------------------------------------------------------------
def build_tables(col: Collector):
    inc = col.final("income_statement")
    years = sorted({k[0] for k in inc})
    income = pd.DataFrame([{"fiscal_year": fy, **{c: inc.get((fy, c)) for c in INCOME_ORDER}} for fy in years])

    cat_d = col.final("net_sales_by_category")
    cat = pd.DataFrame([{"fiscal_year": k[0], "category": k[1], "net_sales": v} for k, v in sorted(cat_d.items())])

    seg_d = col.final("segment_results")
    segs = sorted({(k[0], k[1]) for k in seg_d})
    seg = pd.DataFrame([{"fiscal_year": fy, "region": rg, "net_sales": seg_d.get((fy, rg, "net_sales")),
                         "operating_income": seg_d.get((fy, rg, "operating_income"))} for fy, rg in segs])

    cty_d = col.final("net_sales_by_country")
    cty = pd.DataFrame([{"fiscal_year": k[0], "country": k[1], "net_sales": v} for k, v in sorted(cty_d.items())])

    lines = []
    for (stmt, fy, sec, label), v in sorted(col.final("statement_lines").items(), key=lambda x: (x[0][0], x[0][1])):
        unit = "USD per share" if "per share" in label.lower() else ("thousands of shares" if "(in shares)" in label.lower() else "USD millions")
        lines.append({"statement": stmt, "fiscal_year": fy, "section": sec, "line_item": label, "value": v, "unit": unit})
    return income, cat, seg, cty, pd.DataFrame(lines)


def run_checks(income, cat, seg, cty):
    rows = []

    def add(fy, name, expected, actual):
        if expected is None or actual is None or pd.isna(expected) or pd.isna(actual):
            rows.append({"fiscal_year": fy, "check_name": name, "expected": None, "actual": None, "status": "not enough data"})
        else:
            rows.append({"fiscal_year": fy, "check_name": name, "expected": float(expected), "actual": float(actual),
                         "status": "OK" if abs(expected - actual) <= 1.0 else "CHECK"})

    for r in income.itertuples(index=False):
        fy = r.fiscal_year
        add(fy, "products + services = net sales", r.net_sales,
            None if pd.isna(r.products_net_sales) or pd.isna(r.services_net_sales) else r.products_net_sales + r.services_net_sales)
        add(fy, "net sales - cost of sales = gross margin", r.gross_margin,
            None if pd.isna(r.net_sales) or pd.isna(r.cost_of_sales) else r.net_sales - r.cost_of_sales)
        add(fy, "gross margin - operating expenses = operating income", r.operating_income,
            None if pd.isna(r.gross_margin) or pd.isna(r.total_operating_expenses) else r.gross_margin - r.total_operating_expenses)
        add(fy, "income before taxes - tax = net income", r.net_income,
            None if pd.isna(r.income_before_taxes) or pd.isna(r.income_tax_provision) else r.income_before_taxes - r.income_tax_provision)
        for label, df, col in (("5 product categories add up to net sales", cat, "net_sales"),
                               ("5 regions add up to net sales", seg, "net_sales"),
                               ("US + China + other countries add up to net sales", cty, "net_sales")):
            part = df[df.fiscal_year == fy][col] if not df.empty else []
            add(fy, label, r.net_sales, float(part.sum()) if len(part) else None)
    return pd.DataFrame(rows)


DICTIONARY = [
    ("income_statement", "fiscal_year", "Apple's fiscal year (ends in late September). FY2022 = year ended 24 Sep 2022.", "year"),
    ("income_statement", "net_sales", "Total revenue (Apple calls it net sales).", "USD millions"),
    ("income_statement", "products_net_sales", "Revenue from physical products (iPhone, Mac, iPad, Wearables).", "USD millions"),
    ("income_statement", "services_net_sales", "Revenue from services (App Store, advertising, iCloud, AppleCare ...).", "USD millions"),
    ("income_statement", "cost_of_sales", "Direct cost of producing what was sold.", "USD millions"),
    ("income_statement", "products_cost_of_sales", "Cost of sales for products.", "USD millions"),
    ("income_statement", "services_cost_of_sales", "Cost of sales for services.", "USD millions"),
    ("income_statement", "gross_margin", "Net sales minus cost of sales (profit before running costs).", "USD millions"),
    ("income_statement", "research_and_development", "Spending on developing new products.", "USD millions"),
    ("income_statement", "selling_general_admin", "Sales, marketing and administrative costs.", "USD millions"),
    ("income_statement", "total_operating_expenses", "R&D plus selling, general and administrative.", "USD millions"),
    ("income_statement", "operating_income", "Gross margin minus operating expenses.", "USD millions"),
    ("income_statement", "other_income_expense_net", "Interest, investment gains and other items outside operations.", "USD millions"),
    ("income_statement", "income_before_taxes", "Operating income plus other income/(expense).", "USD millions"),
    ("income_statement", "income_tax_provision", "Income tax charge for the year.", "USD millions"),
    ("income_statement", "net_income", "Profit after tax.", "USD millions"),
    ("income_statement", "eps_basic", "Net income per share (basic).", "USD per share"),
    ("income_statement", "eps_diluted", "Net income per share (diluted: counts share awards).", "USD per share"),
    ("income_statement", "shares_basic_thousands", "Weighted average shares used for basic EPS.", "thousands of shares"),
    ("income_statement", "shares_diluted_thousands", "Weighted average shares used for diluted EPS.", "thousands of shares"),
    ("net_sales_by_category", "category", "Product line: iPhone, Mac, iPad, Wearables/Home/Accessories, Services.", "text"),
    ("net_sales_by_category", "net_sales", "Revenue of that category in the fiscal year.", "USD millions"),
    ("segment_results", "region", "Apple's geographic segment: Americas, Europe, Greater China, Japan, Rest of Asia Pacific.", "text"),
    ("segment_results", "net_sales", "Revenue of the region.", "USD millions"),
    ("segment_results", "operating_income", "Operating profit of the region (before corporate costs).", "USD millions"),
    ("net_sales_by_country", "country", "United States, China (mainland China, Hong Kong, Taiwan) or other countries.", "text"),
    ("net_sales_by_country", "net_sales", "Revenue in that country group.", "USD millions"),
    ("statement_lines", "statement", "Income statement, Balance sheet or Cash flow statement.", "text"),
    ("statement_lines", "section", "Heading the line sits under (for example 'Current assets').", "text"),
    ("statement_lines", "line_item", "The line exactly as Apple labels it.", "text"),
    ("statement_lines", "value", "Amount; see the unit column.", "see unit"),
    ("quality_checks", "check_name", "An accounting rule the data should satisfy, or the item two reports disagree on.", "text"),
    ("quality_checks", "status", "OK = rule holds (within 1 million); CHECK = look at it.", "text"),
]

SAMPLE_QUERIES = [
    ("Annual summary", "SELECT fiscal_year, net_sales, gross_margin_pct, operating_income, net_income, eps_diluted, services_share_pct FROM annual_summary"),
    ("Year-on-year growth by category (SQL window function)",
     "SELECT category, fiscal_year, net_sales, ROUND(100.0*(net_sales - LAG(net_sales) OVER (PARTITION BY category ORDER BY fiscal_year))"
     " / LAG(net_sales) OVER (PARTITION BY category ORDER BY fiscal_year), 1) AS yoy_pct FROM net_sales_by_category"
     " WHERE category IN ('iPhone','Services') ORDER BY category, fiscal_year"),
    ("Greater China share of net sales",
     "SELECT s.fiscal_year, s.net_sales AS china_sales, ROUND(100.0*s.net_sales/i.net_sales,1) AS share_pct"
     " FROM segment_results s JOIN income_statement i USING (fiscal_year) WHERE s.region='Greater China' ORDER BY s.fiscal_year"),
]


def main():
    ap = argparse.ArgumentParser(description="Build Apple's annual SQL database from SEC Financial_Report Excel files.")
    ap.add_argument("--excel", default="excel", help="Folder with the Financial_Report_*.xlsx files (sub-folders are searched too)")
    ap.add_argument("--out", default="database", help="Output folder")
    args = ap.parse_args()

    files = sorted(p for p in Path(args.excel).rglob("*.xlsx") if not p.name.startswith("~$") and "Financial_Report" in p.name)
    if not files:
        sys.exit("No Financial_Report*.xlsx files found in %s" % Path(args.excel).resolve())
    print("Found %d Excel files" % len(files))
    col, log, src_rows = Collector(), [], []
    for p in files:
        res = process_file(p, col, log)
        if res:
            src, years = res
            src_rows.append({"file": p.name, "report_fiscal_year": src, "fiscal_years_covered": ", ".join(map(str, years))})
            print("  read %-34s report FY%d  covers FY%s" % (p.name, src, ", FY".join(map(str, years))))

    income, cat, seg, cty, lines = build_tables(col)
    checks = run_checks(income, cat, seg, cty)
    dis = pd.DataFrame(col.disagreements())
    if dis.empty:
        dis = pd.DataFrame([{"table": "-", "item": "-", "values_by_report": "No disagreements between overlapping reports"}])
    quality = pd.concat([checks.assign(kind="accounting rule"),
                         dis.rename(columns={"item": "check_name", "values_by_report": "detail"}).assign(kind="reports disagree")],
                        ignore_index=True)
    dictionary = pd.DataFrame(DICTIONARY, columns=["table", "column", "meaning", "unit"])
    sources = pd.DataFrame(src_rows)

    out = Path(args.out)
    (out / "tables").mkdir(parents=True, exist_ok=True)
    tables = {"income_statement": income, "net_sales_by_category": cat, "segment_results": seg,
              "net_sales_by_country": cty, "statement_lines": lines, "quality_checks": quality,
              "data_dictionary": dictionary, "source_files": sources}
    db = out / "apple.db"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    for name, df in tables.items():
        df.to_sql(name, con, index=False)
        df.to_csv(out / "tables" / (name + ".csv"), index=False)
    con.execute("""CREATE VIEW annual_summary AS
        SELECT fiscal_year, net_sales, gross_margin, ROUND(100.0*gross_margin/net_sales,1) AS gross_margin_pct,
               operating_income, net_income, eps_diluted, services_net_sales,
               ROUND(100.0*services_net_sales/net_sales,1) AS services_share_pct
        FROM income_statement ORDER BY fiscal_year""")
    con.commit()
    with pd.ExcelWriter(out / "Apple_Annual_Tables.xlsx", engine="openpyxl") as w:
        for name, df in tables.items():
            df.to_excel(w, sheet_name=name[:31], index=False)

    print("\nTables written to %s" % db.resolve())
    for name, df in tables.items():
        print("  %-24s %5d rows" % (name, len(df)))
    bad = checks[checks.status == "CHECK"]
    print("\nAccounting checks: %d OK, %d CHECK, %d not enough data" % (
        (checks.status == "OK").sum(), len(bad), (checks.status == "not enough data").sum()))
    if len(bad):
        print(bad.to_string(index=False))
    if "No disagreements" not in str(dis.iloc[0].get("values_by_report", "")):
        print("\nReports disagree on %d values - see quality_checks (kind = 'reports disagree')" % len(dis))
    else:
        print("Overlapping reports agree with each other.")
    for line in log:
        print(line)

    print("\n--- Try these SQL queries ---")
    for title, sql in SAMPLE_QUERIES:
        print("\n# %s\n%s" % (title, sql))
        print(pd.read_sql_query(sql, con).to_string(index=False))
    con.close()


if __name__ == "__main__":
    main()
