#!/usr/bin/env python3
"""
build_final_tables.py
=====================
Combines your yearly SEC "Financial_Report" files (the CSV folders, or the .xlsx files)
into ONE clean, final set of CSV tables for DuckDB.

Put this file in your  ai_assistant  folder (next to build_database.py) and run:

    python build_final_tables.py

It reads:   tables/Financial_Report_2021_CSV/  ... Financial_Report_2024_CSV/   (folders of CSV files)
            tables/Financial_Report_2025.xlsx                                    (also accepted)
It writes:  final_data/
              income_statement.csv       one row per fiscal year
              balance_sheet.csv          one row per fiscal year
              cash_flow.csv              one row per fiscal year
              sales_by_category.csv      iPhone / Mac / iPad / Wearables / Services
              sales_by_region.csv        Americas / Europe / Greater China / Japan / Rest of Asia Pacific
              sales_by_country.csv       U.S. / China / Other countries
              all_facts_long.csv         every number found, one per row (nothing is lost)
              data_checks.csv            does everything add up? did any year change between reports?
              data_dictionary.csv        what every column means (useful for text-to-SQL later)
              apple_final.duckdb         the same tables, loaded into a DuckDB database
              load_into_duckdb.sql       the SQL that loads the CSVs, if you prefer to do it yourself

It never changes or deletes your existing files.
All money amounts are USD millions unless the column name says otherwise.
"""
from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

# ----------------------------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------------------------
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
DATE_TXT = re.compile(r"^([A-Za-z]{3,9})\.?\s+(\d{1,2}),\s*(\d{4})$")
DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


def parse_date(c):
    if isinstance(c, datetime):
        return c.date()
    if isinstance(c, date):
        return c
    if not isinstance(c, str):
        return None
    s = c.strip()
    m = DATE_TXT.match(s)
    if m and m.group(1)[:3].lower() in MONTHS:
        return date(int(m.group(3)), MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
    m = DATE_ISO.match(s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def norm_cell(c):
    """Turn any cell into None / text / date / number-as-is."""
    if c is None:
        return None
    if isinstance(c, (datetime, date, int, float)) and not isinstance(c, bool):
        return c
    s = str(c).replace("\xa0", " ").strip()
    return s or None


def to_num(c):
    """'365,817' -> 365817.0 ; '(1,234)' -> -1234.0 ; '$ 5.67' -> 5.67 ; text -> None"""
    if isinstance(c, bool) or c is None:
        return None
    if isinstance(c, (int, float)):
        return float(c)
    if not isinstance(c, str):
        return None
    s = c.replace("$", "").replace(" ", "").strip()
    if not s:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    s = s.replace(",", "")
    if s.startswith("-"):
        neg, s = True, s[1:]
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def fmt(v):
    if v is None:
        return ""
    if abs(v - round(v)) < 1e-9:
        return str(int(round(v)))
    return repr(round(v, 4))


# ----------------------------------------------------------------------------------------------
# Reading a report (CSV folder or xlsx file)  ->  list of "tables" (one per sheet)
# ----------------------------------------------------------------------------------------------
def read_csv_rows(path: Path):
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    return [[norm_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]


def read_xlsx_sheets(path: Path):
    from openpyxl import load_workbook  # only needed for .xlsx input

    wb = load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        yield name, [[norm_cell(c) for c in row] for row in wb[name].iter_rows(values_only=True)]


class Table:
    def __init__(self, source, title, cols, body):
        self.source, self.title, self.cols, self.body = source, title, cols, body


def parse_table(source, rows):
    """Find the title and the date columns, no matter how the sheet was exported."""
    title, ti = None, None
    for i, r in enumerate(rows[:5]):
        c0 = r[0] if r else None
        if isinstance(c0, str) and len(c0) > 8 and re.search(r"[A-Za-z]{3}", c0) and not parse_date(c0):
            title, ti = c0, i
            break
    if title is None:
        return None
    for j in range(ti, min(ti + 4, len(rows))):
        cols = [(k, parse_date(c)) for k, c in enumerate(rows[j]) if k > 0 and parse_date(c)]
        if cols:
            return Table(source, title, cols, rows[j + 1:])
    return None  # e.g. cover page, notes without numbers by date


def kind_of(title: str):
    t = title.lower()
    if t.startswith("consolidated statements of operations"):
        return "income_statement"
    if t.startswith("consolidated balance sheets") and "parenthetical" not in t:
        return "balance_sheet"
    if t.startswith("consolidated statements of cash flows"):
        return "cash_flow"
    if "net sales disaggregated by significant products and services" in t:
        return "category"
    if "information by reportable segment" in t and "(details)" in t:
        return "region"
    if "segment information and geographic data" in t and "net sales (details)" in t:
        return "country"
    return None


def lines(table):
    for r in table.body:
        label = r[0] if r else None
        if not isinstance(label, str):
            continue
        vals = [to_num(r[k]) if k < len(r) else None for k, _ in table.cols]
        yield label.strip(), vals


# ----------------------------------------------------------------------------------------------
# Section tracking and mapping of line items to clean column names
# ----------------------------------------------------------------------------------------------
def header_key(label):
    return label.strip().rstrip(":").strip().lower().replace("’", "'")


def next_section(kind, prev, label):
    h = header_key(label)
    if kind == "income_statement":
        if h in ("net sales", "cost of sales", "products", "services"):
            return h.replace(" ", "_")
        if h == "earnings per share":
            return "eps"
        if h.startswith("shares used in computing"):
            return "shares"
        if h == "operating expenses":
            return "opex"
        return prev
    return h


def map_income_statement(sec, label):
    l = label.lower().strip()
    if l == "total net sales":
        return "net_sales"
    if l == "total cost of sales":
        return "cost_of_sales"
    if l == "net sales":
        return {"products": "products_net_sales", "services": "services_net_sales"}.get(sec, "net_sales")
    if l == "cost of sales":
        return {"products": "products_cost_of_sales", "services": "services_cost_of_sales"}.get(sec, "cost_of_sales")
    if l in ("products", "services"):
        if sec == "net_sales":
            return l + "_net_sales"
        if sec == "cost_of_sales":
            return l + "_cost_of_sales"
    if l == "gross margin":
        return "gross_margin"
    if l.startswith("research and development"):
        return "research_and_development"
    if l.startswith("selling, general"):
        return "selling_general_admin"
    if l.startswith("total operating expenses"):
        return "total_operating_expenses"
    if l.startswith("operating income"):
        return "operating_income"
    if l.startswith("other income"):
        return "other_income_expense_net"
    if l.startswith("income before provision"):
        return "income_before_tax"
    if l.startswith("provision for income taxes"):
        return "income_tax_provision"
    if l == "net income":
        return "net_income"
    if l.startswith("basic") or l.startswith("diluted"):
        which = "basic" if l.startswith("basic") else "diluted"
        if "per share" in l or sec == "eps":
            return "eps_" + which
        if "in shares" in l or sec == "shares":
            return "shares_%s_thousands" % which
    return None


def map_balance_sheet(sec, label):
    l = label.lower().replace("’", "'").strip()
    non = sec.startswith("non-current")
    if l.startswith("cash and cash equivalents"):
        return "cash_and_equivalents"
    if l.startswith("marketable securities"):
        return "marketable_securities_non_current" if non else "marketable_securities_current"
    if l.startswith("accounts receivable"):
        return "accounts_receivable_net"
    if l.startswith("inventories"):
        return "inventories"
    if l.startswith("vendor non-trade receivables"):
        return "vendor_non_trade_receivables"
    if l == "other current assets":
        return "other_current_assets"
    if l == "total current assets":
        return "total_current_assets"
    if l.startswith("property, plant and equipment"):
        return "property_plant_equipment_net"
    if l == "other non-current assets":
        return "other_non_current_assets"
    if l == "total non-current assets":
        return "total_non_current_assets"
    if l == "total assets":
        return "total_assets"
    if l.startswith("accounts payable"):
        return "accounts_payable"
    if l == "other current liabilities":
        return "other_current_liabilities"
    if l.startswith("deferred revenue"):
        return "deferred_revenue"
    if l.startswith("commercial paper"):
        return "commercial_paper"
    if l.startswith("term debt"):
        return "term_debt_non_current" if non else "term_debt_current"
    if l == "other non-current liabilities":
        return "other_non_current_liabilities"
    if l == "total current liabilities":
        return "total_current_liabilities"
    if l == "total non-current liabilities":
        return "total_non_current_liabilities"
    if l == "total liabilities":
        return "total_liabilities"
    if l.startswith("common stock and additional paid-in capital"):
        return "common_stock_and_apic"
    if l.startswith("retained earnings") or l.startswith("accumulated deficit"):
        return "retained_earnings_or_deficit"
    if l.startswith("accumulated other comprehensive"):
        return "accumulated_other_comprehensive"
    if l.startswith("total liabilities and shareholders"):
        return "total_liabilities_and_equity"
    if l.startswith("total shareholders"):
        return "total_shareholders_equity"
    return None


def map_cash_flow(sec, label):
    l = label.lower().strip()
    if l.startswith("cash, cash equivalents and restricted cash, beginning"):
        return "cash_beginning"
    if l.startswith("cash, cash equivalents and restricted cash, ending"):
        return "cash_ending"
    if l == "net income":
        return "net_income"
    if l == "depreciation and amortization":
        return "depreciation_amortization"
    if l.startswith("share-based compensation"):
        return "share_based_compensation"
    if re.match(r"^cash.*operating activities$", l):
        return "operating_cash_flow"
    if re.match(r"^cash.*investing activities$", l):
        return "investing_cash_flow"
    if re.match(r"^cash.*financing activities$", l):
        return "financing_cash_flow"
    if l.startswith("payments for acquisition of property, plant and equipment"):
        return "capital_expenditures"
    if l.startswith("payments for dividends"):
        return "dividends_paid"
    if l.startswith("repurchases of common stock"):
        return "share_repurchases"
    if re.match(r"^(increase|decrease)[^,]* in cash", l):
        return "net_change_in_cash"
    return None


MAPPERS = {"income_statement": map_income_statement, "balance_sheet": map_balance_sheet, "cash_flow": map_cash_flow}

IS_COLS = ["net_sales", "products_net_sales", "services_net_sales", "cost_of_sales", "products_cost_of_sales",
           "services_cost_of_sales", "gross_margin", "research_and_development", "selling_general_admin",
           "total_operating_expenses", "operating_income", "other_income_expense_net", "income_before_tax",
           "income_tax_provision", "net_income", "eps_basic", "eps_diluted", "shares_basic_thousands",
           "shares_diluted_thousands"]
BS_COLS = ["cash_and_equivalents", "marketable_securities_current", "accounts_receivable_net", "inventories",
           "vendor_non_trade_receivables", "other_current_assets", "total_current_assets",
           "marketable_securities_non_current", "property_plant_equipment_net", "other_non_current_assets",
           "total_non_current_assets", "total_assets", "accounts_payable", "other_current_liabilities",
           "deferred_revenue", "commercial_paper", "term_debt_current", "total_current_liabilities",
           "term_debt_non_current", "other_non_current_liabilities", "total_non_current_liabilities",
           "total_liabilities", "common_stock_and_apic", "retained_earnings_or_deficit",
           "accumulated_other_comprehensive", "total_shareholders_equity", "total_liabilities_and_equity"]
CF_COLS = ["net_income", "depreciation_amortization", "share_based_compensation", "operating_cash_flow",
           "capital_expenditures", "investing_cash_flow", "dividends_paid", "share_repurchases",
           "financing_cash_flow", "net_change_in_cash", "cash_beginning", "cash_ending"]
WIDE = {"income_statement": IS_COLS, "balance_sheet": BS_COLS, "cash_flow": CF_COLS}

CATEGORY_ORDER = ["iPhone", "Mac", "iPad", "Wearables, Home and Accessories", "Services"]
REGION_ORDER = ["Americas", "Europe", "Greater China", "Japan", "Rest of Asia Pacific"]
COUNTRY_ORDER = ["U.S.", "China", "Other countries"]

DICTIONARY = {
    "income_statement": {
        "fiscal_year": ("Apple's fiscal year (ends in late September). FY2024 = year ended 28 Sep 2024.", ""),
        "period_end": ("Last day of the fiscal year.", "date"),
        "net_sales": ("Total net sales (revenue).", "USD millions"),
        "products_net_sales": ("Net sales of products (iPhone, Mac, iPad, Wearables).", "USD millions"),
        "services_net_sales": ("Net sales of Services.", "USD millions"),
        "cost_of_sales": ("Total cost of sales.", "USD millions"),
        "products_cost_of_sales": ("Cost of sales for products.", "USD millions"),
        "services_cost_of_sales": ("Cost of sales for services.", "USD millions"),
        "gross_margin": ("Net sales minus cost of sales (profit before operating expenses).", "USD millions"),
        "research_and_development": ("Research and development expense.", "USD millions"),
        "selling_general_admin": ("Selling, general and administrative expense.", "USD millions"),
        "total_operating_expenses": ("R&D plus SG&A.", "USD millions"),
        "operating_income": ("Gross margin minus operating expenses.", "USD millions"),
        "other_income_expense_net": ("Other income/(expense), net.", "USD millions"),
        "income_before_tax": ("Income before provision for income taxes.", "USD millions"),
        "income_tax_provision": ("Provision for income taxes.", "USD millions"),
        "net_income": ("Profit after tax.", "USD millions"),
        "eps_basic": ("Basic earnings per share.", "USD per share"),
        "eps_diluted": ("Diluted earnings per share.", "USD per share"),
        "shares_basic_thousands": ("Weighted average shares used for basic EPS.", "thousands of shares"),
        "shares_diluted_thousands": ("Weighted average shares used for diluted EPS.", "thousands of shares"),
        "source_report": ("Which Financial_Report year this row was taken from (the latest report that contains this year).", ""),
    },
    "balance_sheet": {
        "fiscal_year": ("Apple's fiscal year.", ""),
        "period_end": ("Balance sheet date.", "date"),
        "cash_and_equivalents": ("Cash and cash equivalents.", "USD millions"),
        "marketable_securities_current": ("Marketable securities held as current assets.", "USD millions"),
        "accounts_receivable_net": ("Accounts receivable, net.", "USD millions"),
        "inventories": ("Inventories.", "USD millions"),
        "vendor_non_trade_receivables": ("Vendor non-trade receivables.", "USD millions"),
        "other_current_assets": ("Other current assets.", "USD millions"),
        "total_current_assets": ("Total current assets.", "USD millions"),
        "marketable_securities_non_current": ("Marketable securities held long term.", "USD millions"),
        "property_plant_equipment_net": ("Property, plant and equipment, net.", "USD millions"),
        "other_non_current_assets": ("Other non-current assets.", "USD millions"),
        "total_non_current_assets": ("Total non-current assets.", "USD millions"),
        "total_assets": ("Total assets.", "USD millions"),
        "accounts_payable": ("Accounts payable.", "USD millions"),
        "other_current_liabilities": ("Other current liabilities.", "USD millions"),
        "deferred_revenue": ("Deferred revenue (current).", "USD millions"),
        "commercial_paper": ("Commercial paper (short-term borrowing).", "USD millions"),
        "term_debt_current": ("Term debt due within one year.", "USD millions"),
        "total_current_liabilities": ("Total current liabilities.", "USD millions"),
        "term_debt_non_current": ("Term debt due after one year.", "USD millions"),
        "other_non_current_liabilities": ("Other non-current liabilities.", "USD millions"),
        "total_non_current_liabilities": ("Total non-current liabilities.", "USD millions"),
        "total_liabilities": ("Total liabilities.", "USD millions"),
        "common_stock_and_apic": ("Common stock and additional paid-in capital.", "USD millions"),
        "retained_earnings_or_deficit": ("Retained earnings (negative = accumulated deficit).", "USD millions"),
        "accumulated_other_comprehensive": ("Accumulated other comprehensive income/(loss).", "USD millions"),
        "total_shareholders_equity": ("Total shareholders' equity.", "USD millions"),
        "total_liabilities_and_equity": ("Total liabilities and shareholders' equity (equals total assets).", "USD millions"),
        "source_report": ("Which Financial_Report year this row was taken from.", ""),
    },
    "cash_flow": {
        "fiscal_year": ("Apple's fiscal year.", ""),
        "period_end": ("Last day of the fiscal year.", "date"),
        "net_income": ("Net income (start of the cash flow statement).", "USD millions"),
        "depreciation_amortization": ("Depreciation and amortization.", "USD millions"),
        "share_based_compensation": ("Share-based compensation expense.", "USD millions"),
        "operating_cash_flow": ("Cash generated by operating activities.", "USD millions"),
        "capital_expenditures": ("Payments for property, plant and equipment (negative = cash out).", "USD millions"),
        "investing_cash_flow": ("Cash generated by/(used in) investing activities.", "USD millions"),
        "dividends_paid": ("Payments for dividends and dividend equivalents (negative = cash out).", "USD millions"),
        "share_repurchases": ("Repurchases of common stock (negative = cash out).", "USD millions"),
        "financing_cash_flow": ("Cash used in financing activities.", "USD millions"),
        "net_change_in_cash": ("Increase/(decrease) in cash, cash equivalents and restricted cash.", "USD millions"),
        "cash_beginning": ("Cash, cash equivalents and restricted cash, beginning of year.", "USD millions"),
        "cash_ending": ("Cash, cash equivalents and restricted cash, end of year.", "USD millions"),
        "source_report": ("Which Financial_Report year this row was taken from.", ""),
    },
    "sales_by_category": {
        "fiscal_year": ("Apple's fiscal year.", ""), "period_end": ("Last day of the fiscal year.", "date"),
        "category": ("iPhone, Mac, iPad, Wearables, Home and Accessories, or Services.", ""),
        "net_sales_usd_m": ("Net sales of that category.", "USD millions"),
        "source_report": ("Which Financial_Report year this row was taken from.", ""),
    },
    "sales_by_region": {
        "fiscal_year": ("Apple's fiscal year.", ""), "period_end": ("Last day of the fiscal year.", "date"),
        "region": ("Apple's reportable segment (geographic region).", ""),
        "net_sales_usd_m": ("Net sales of that region.", "USD millions"),
        "operating_income_usd_m": ("Segment operating income of that region (before R&D and corporate costs).", "USD millions"),
        "source_report": ("Which Financial_Report year this row was taken from.", ""),
    },
    "sales_by_country": {
        "fiscal_year": ("Apple's fiscal year.", ""), "period_end": ("Last day of the fiscal year.", "date"),
        "country_group": ("U.S., China, or Other countries.", ""),
        "net_sales_usd_m": ("Net sales for that country group.", "USD millions"),
        "source_report": ("Which Financial_Report year this row was taken from.", ""),
    },
}


def unit_for(kind, sec, label):
    l = label.lower()
    if kind == "income_statement":
        if "per share" in l or sec == "eps":
            return "USD per share"
        if "in shares" in l or sec == "shares":
            return "thousands of shares"
    return "USD millions"


# ----------------------------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------------------------
def find_reports(tables_dir: Path):
    reports = {}  # year -> ("csv", folder) or ("xlsx", file)
    for p in sorted(tables_dir.iterdir()):
        m = re.match(r"Financial_Report_(\d{4})", p.name, re.I)
        if not m:
            continue
        year = int(m.group(1))
        if p.is_dir():
            reports[year] = ("csv", p)  # CSV folder wins if both exist
        elif p.suffix.lower() == ".xlsx":
            reports.setdefault(year, ("xlsx", p))
    return reports


def main():
    try:  # keep Windows consoles from crashing on unusual characters
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Combine yearly Financial_Report files into final DuckDB-ready tables.")
    ap.add_argument("--tables", default=str(here / "tables"), help="Folder that holds the Financial_Report_YYYY folders/files")
    ap.add_argument("--out", default=str(here / "final_data"), help="Output folder (created if missing)")
    ap.add_argument("--no-duckdb", action="store_true", help="Only write CSV files, skip the .duckdb database")
    args = ap.parse_args()

    tables_dir, out = Path(args.tables), Path(args.out)
    if not tables_dir.is_dir():
        sys.exit("I cannot find the tables folder: %s\nRun this script from your ai_assistant folder, or use --tables <folder>." % tables_dir)
    reports = find_reports(tables_dir)
    if not reports:
        sys.exit("No Financial_Report_YYYY folders or .xlsx files found in %s" % tables_dir)
    out.mkdir(parents=True, exist_ok=True)

    print("Reading yearly reports from: %s" % tables_dir)
    LONG = defaultdict(dict)  # (kind, section, label, fy) -> {report_year: (value, period_end)}
    for ry in sorted(reports):
        how, path = reports[ry]
        if how == "csv":
            sheets = ((f.name, read_csv_rows(f)) for f in sorted(path.rglob("*.csv")))
        else:
            sheets = read_xlsx_sheets(path)
        found = set()
        for name, rows in sheets:
            t = parse_table(name, rows)
            if not t:
                continue
            kind = kind_of(t.title)
            if not kind:
                continue
            found.add(kind)
            sec, seen = "", defaultdict(int)
            for label, vals in lines(t):
                if all(v is None for v in vals):  # a heading line
                    if kind in ("category", "region", "country"):
                        if not label.endswith("]"):  # skip "[Line Items]" lines, keep member names
                            sec = label
                    else:
                        sec = next_section(kind, sec, label)
                    continue
                seen[(sec, label)] += 1
                lab = label if seen[(sec, label)] == 1 else "%s #%d" % (label, seen[(sec, label)])
                for (k, pe), v in zip(t.cols, vals):
                    if v is None:
                        continue
                    LONG[(kind, sec, lab, pe.year)][ry] = (v, pe)
        want = ["income_statement", "balance_sheet", "cash_flow", "category", "region", "country"]
        missing = [w for w in want if w not in found]
        print("  Financial_Report_%d  found: %d of 6 tables%s" % (ry, 6 - len(missing),
              ("   (NOT found: %s)" % ", ".join(missing)) if missing else ""))

    # choose the newest report for every number and note where reports disagree
    picked, disagreements = {}, []
    for key, per in LONG.items():
        latest = max(per)
        val, pe = per[latest]
        agree = all(abs(v - val) <= (0.005 if key[0] == "income_statement" and "per share" in key[2].lower() else 0.5) for v, _ in per.values())
        picked[key] = (val, pe, latest, len(per), agree)
        if not agree:
            for ry, (v, _) in sorted(per.items()):
                if ry != latest:
                    disagreements.append((key, ry, v, latest, val))

    # ---- statement tables (one row per fiscal year)
    wide, unmapped, collisions = {k: defaultdict(dict) for k in WIDE}, set(), []
    for (kind, sec, label, fy), (val, pe, ry, n, agree) in picked.items():
        if kind not in WIDE:
            continue
        col = MAPPERS[kind](sec, label)
        if col is None:
            unmapped.add((kind, sec, label))
            continue
        row = wide[kind][fy]
        if col in row and abs(row[col] - val) > 1e-9:
            collisions.append((kind, fy, col, label))
            continue
        row[col] = val
        row["period_end"] = pe
        row["_ry"] = max(row.get("_ry", 0), ry)

    def write_csv(name, header, rows):
        with open(out / (name + ".csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)

    table_rows = {}
    for kind, cols in WIDE.items():
        rows = []
        for fy in sorted(wide[kind]):
            r = wide[kind][fy]
            rows.append([fy, r["period_end"].isoformat()] + [fmt(r.get(c)) for c in cols] + [r["_ry"]])
        write_csv(kind, ["fiscal_year", "period_end"] + cols + ["source_report"], rows)
        table_rows[kind] = rows

    # ---- members: category / region / country
    def members(kind, field):
        d = defaultdict(dict)
        for (k, sec, label, fy), (val, pe, ry, n, agree) in picked.items():
            if k == kind and sec and label == field:
                d[(fy, sec)] = (val, pe, ry)
        return d

    cat = members("category", "Net sales")
    reg_s, reg_o = members("region", "Net sales"), members("region", "Operating income")
    cou = members("country", "Net sales")

    def ordered(d, order):
        extra = sorted({m for (_, m) in d} - set(order))
        return order + extra, extra

    cat_order, cat_extra = ordered(cat, CATEGORY_ORDER)
    rows = [[fy, cat[(fy, m)][1].isoformat(), m, fmt(cat[(fy, m)][0]), cat[(fy, m)][2]]
            for fy in sorted({k[0] for k in cat}) for m in cat_order if (fy, m) in cat]
    write_csv("sales_by_category", ["fiscal_year", "period_end", "category", "net_sales_usd_m", "source_report"], rows)

    reg_order, reg_extra = ordered(reg_s, REGION_ORDER)
    rows = []
    for fy in sorted({k[0] for k in reg_s}):
        for m in reg_order:
            if (fy, m) in reg_s:
                v, pe, ry = reg_s[(fy, m)]
                o = reg_o.get((fy, m))
                rows.append([fy, pe.isoformat(), m, fmt(v), fmt(o[0]) if o else "", ry])
    write_csv("sales_by_region", ["fiscal_year", "period_end", "region", "net_sales_usd_m", "operating_income_usd_m", "source_report"], rows)

    cou_order, cou_extra = ordered(cou, COUNTRY_ORDER)
    rows = [[fy, cou[(fy, m)][1].isoformat(), m, fmt(cou[(fy, m)][0]), cou[(fy, m)][2]]
            for fy in sorted({k[0] for k in cou}) for m in cou_order if (fy, m) in cou]
    write_csv("sales_by_country", ["fiscal_year", "period_end", "country_group", "net_sales_usd_m", "source_report"], rows)

    # ---- everything, long
    rows = []
    for (kind, sec, label, fy), (val, pe, ry, n, agree) in sorted(picked.items(), key=lambda x: (x[0][0], x[0][3], x[0][1], x[0][2])):
        rows.append([kind, sec, label, fy, pe.isoformat(), fmt(val), unit_for(kind, sec, label), ry, n, "yes" if agree else "NO"])
    write_csv("all_facts_long", ["table_name", "section_or_member", "line_item", "fiscal_year", "period_end", "value",
                                 "unit", "source_report", "reports_containing_it", "same_in_all_reports"], rows)

    # ---- checks
    checks = []

    def chk(name, fy, expected, actual):
        if expected is None or actual is None:
            return
        d = actual - expected
        status = "OK" if abs(d) < 0.5 else ("OK_rounding" if abs(d) <= 2 else "CHECK")
        checks.append([name, fy, fmt(expected), fmt(actual), fmt(d), status])

    g = lambda kind, fy, c: wide[kind].get(fy, {}).get(c)
    for fy in sorted(wide["income_statement"]):
        ns, cs, gm = (g("income_statement", fy, c) for c in ("net_sales", "cost_of_sales", "gross_margin"))
        chk("gross_margin = net_sales - cost_of_sales", fy, gm, ns - cs if ns is not None and cs is not None else None)
        og, ox, oi = (g("income_statement", fy, c) for c in ("gross_margin", "total_operating_expenses", "operating_income"))
        chk("operating_income = gross_margin - operating_expenses", fy, oi, og - ox if og is not None and ox is not None else None)
        ib, tx, ni = (g("income_statement", fy, c) for c in ("income_before_tax", "income_tax_provision", "net_income"))
        chk("net_income = income_before_tax - income_tax", fy, ni, ib - tx if ib is not None and tx is not None else None)
        ps, ss = g("income_statement", fy, "products_net_sales"), g("income_statement", fy, "services_net_sales")
        chk("products + services net sales = net_sales", fy, ns, ps + ss if ps is not None and ss is not None else None)
        cats = [cat[(fy, m)][0] for m in CATEGORY_ORDER if (fy, m) in cat]
        if len(cats) == len(CATEGORY_ORDER):
            chk("sum of 5 categories = net_sales", fy, ns, sum(cats))
        regs = [reg_s[(fy, m)][0] for m in REGION_ORDER if (fy, m) in reg_s]
        if len(regs) == len(REGION_ORDER):
            chk("sum of 5 regions = net_sales", fy, ns, sum(regs))
        cous = [cou[(fy, m)][0] for m in COUNTRY_ORDER if (fy, m) in cou]
        if len(cous) == len(COUNTRY_ORDER):
            chk("sum of U.S. + China + other = net_sales", fy, ns, sum(cous))
    for fy in sorted(wide["balance_sheet"]):
        ta, tl, te, tle = (g("balance_sheet", fy, c) for c in ("total_assets", "total_liabilities", "total_shareholders_equity", "total_liabilities_and_equity"))
        chk("total_assets = total_liabilities + equity", fy, ta, tl + te if tl is not None and te is not None else None)
        chk("total_assets = total_liabilities_and_equity", fy, ta, tle)
        ca, na = g("balance_sheet", fy, "total_current_assets"), g("balance_sheet", fy, "total_non_current_assets")
        chk("current + non-current assets = total_assets", fy, ta, ca + na if ca is not None and na is not None else None)
    for fy in sorted(wide["cash_flow"]):
        o, i, f_, ch = (g("cash_flow", fy, c) for c in ("operating_cash_flow", "investing_cash_flow", "financing_cash_flow", "net_change_in_cash"))
        chk("operating + investing + financing = net change in cash", fy, ch, o + i + f_ if None not in (o, i, f_) else None)
        b, e = g("cash_flow", fy, "cash_beginning"), g("cash_flow", fy, "cash_ending")
        chk("cash_beginning + net change = cash_ending", fy, e, b + ch if b is not None and ch is not None else None)
    for (kind, sec, label, fy), ry, v, latest, val in disagreements:
        checks.append(["same value in every report: %s / %s / %s" % (kind, sec or "-", label), fy, fmt(val), fmt(v),
                       fmt(v - val), "CHANGED between FY%d report and FY%d report (kept FY%d)" % (ry, latest, latest)])
    write_csv("data_checks", ["check", "fiscal_year", "expected", "actual", "difference", "status"], checks)

    # ---- data dictionary
    rows = [[t, c, d, u] for t, cols in DICTIONARY.items() for c, (d, u) in cols.items()]
    write_csv("data_dictionary", ["table_name", "column_name", "description", "unit"], rows)

    # ---- DuckDB
    names = ["income_statement", "balance_sheet", "cash_flow", "sales_by_category", "sales_by_region",
             "sales_by_country", "all_facts_long", "data_checks", "data_dictionary"]
    sql_lines = ["-- Loads the final CSV files into DuckDB. Generated by build_final_tables.py"]
    for n in names:
        p = (out / (n + ".csv")).resolve().as_posix()
        sql_lines.append("CREATE OR REPLACE TABLE %s AS SELECT * FROM read_csv('%s', header = true, auto_detect = true);" % (n, p))
    (out / "load_into_duckdb.sql").write_text("\n".join(sql_lines) + "\n", encoding="utf-8")
    db_msg = ""
    if not args.no_duckdb:
        try:
            import duckdb

            db = out / "apple_final.duckdb"
            con = duckdb.connect(str(db))
            for line in sql_lines[1:]:
                con.execute(line)
            con.close()
            db_msg = "  Database : %s" % db
        except ImportError:
            db_msg = "  (DuckDB is not installed in this Python, so apple_final.duckdb was not created. Run: pip install duckdb)"

    # ---- plain-English report
    bad = [c for c in checks if c[5] == "CHECK"]
    changed = [c for c in checks if str(c[5]).startswith("CHANGED")]
    years = sorted(wide["income_statement"])
    print("\nDONE.")
    print("  Fiscal years found : %s" % (", ".join("FY%d" % y for y in years) if years else "none"))
    print("  Output folder      : %s" % out.resolve())
    for k in ("income_statement", "balance_sheet", "cash_flow"):
        print("    %-18s %d rows" % (k + ".csv", len(table_rows[k])))
    print("    sales_by_category.csv / sales_by_region.csv / sales_by_country.csv / all_facts_long.csv / data_checks.csv / data_dictionary.csv")
    if db_msg:
        print(db_msg)
    n_ok = sum(1 for c in checks if c[5] in ("OK", "OK_rounding"))
    print("\nChecks: %d passed, %d failed, %d numbers changed between reports." % (n_ok, len(bad), len(changed)))
    for c in bad[:15]:
        print("  FAILED: %s  FY%s  expected %s, got %s" % (c[0], c[1], c[2], c[3]))
    if changed:
        print("  (Apple restated some earlier years; the newest report's number is kept. See data_checks.csv.)")
    unmapped = {u for u in unmapped if u[0] != "cash_flow"}  # cash flow detail lines are deliberately kept only in all_facts_long.csv
    if unmapped:
        print("\nNote: %d line items were not given a column name. They are still in all_facts_long.csv:" % len(unmapped))
        for k in sorted(unmapped)[:12]:
            print("   - %s | %s | %s" % k)
    if collisions:
        print("\nNote: %d line items shared a column name with different values (first one kept):" % len(collisions))
        for c in collisions[:8]:
            print("   - %s FY%d %s (%s)" % c)
    extras = cat_extra + reg_extra + cou_extra
    if extras:
        print("\nNote: extra category/region names found (kept in the CSVs): %s" % ", ".join(extras))


if __name__ == "__main__":
    main()
